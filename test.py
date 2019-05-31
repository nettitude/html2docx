from docx import Document
from h2d import HtmlToDocx

document = Document()
table = document.add_table(3,3)
cell = table.cell(1,1)
parser = HtmlToDocx()

text1 = '''<p>My line <span width="10" style="color: rgb(235, 107, 86);">goes he</span>re</p>
<p><span style="background-color: rgb(251, 160, 38);">Background color</span></p>
<p><span style="background-color: rgb(71, 85, 119); color: rgb(255, 255, 255);">This sentence has background
        and&nbsp;</span><span style="background-color: rgb(71, 85, 119);"><span style="color: rgb(247, 218, 100);">text
            color</span><span style="color: rgb(100, 100, 100);">&nbsp;</span></span><span
        style="background-color: rgb(71, 85, 119); color: rgb(255, 255, 255);">and<strong> bold </strong><em>italic</em>
        <u>underlined</u> <s>strike</s> styles</span></p>
<p>List</p>
<h1>heading 1</h1>
<ol>
    <li>A list for reals</li>
    <p>hey</p>
    <li>Second item</li>
</ol>
<ul style="list-style-type: circle;">
    <li>Unorderd list</li>
    <li>with circle markers</li>
</ul>
<li>yo</li>
<p>Align left Align leftAlign leftAlign leftAlign leftAlign leftAlign leftAlign leftAlign leftAlign left</p>
<p style="text-align: center;">Align center Align center Align center Align center Align center Align center Align
    center Align center Align center</p>
<p style="text-align: right;">Align Right. Align Right. Align Right. Align Right. Align Right. Align Right. Align Right.
    Align Right.</p>
<p style="text-align: justify;">This sentence is justified. This sentence is justified. This sentence is justified. This
    sentence is justified. This sentence is justified.</p>
<p style="text-align: justify;">Indent 0</p>
<p style="text-align: justify; margin-left: 20px;">Indent 1</p>
<p style="text-align: justify; margin-left: 40px;">Indent 2</p>
<p style="text-align: justify; margin-left: 60px;">Indent 3</p>
<p style="text-align: justify; margin-left: 80px;">Indent 4</p>
<p style="text-align: justify; margin-left: 580px;">Indent max?</p>
<p style="text-align: left;">asdfsa</p>
<p style="text-align: left;"><a class="fr-green fr-strong" href="https://github.com" rel="noopener noreferrer"
        target="_blank">link</a></p>
<ol>
    <li>hey</li>
    <ul>
        <li>hey2</li>
        <ul>
            <ol>
        </ul>
</ol>
</ol>
</ul>'''

parser.add_html_to_document(text1, document)
parser.add_html_to_document(text1, cell)

document.save('test.docx')